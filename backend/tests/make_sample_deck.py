"""Generate a small sample pitch deck for e2e testing.

Run from backend/:  ../.venv/bin/python tests/make_sample_deck.py
Writes: /tmp path printed on stdout.
"""
import tempfile
from pathlib import Path

from pptx import Presentation

SLIDES = [
    ("Rover — dog walking, on demand",
     "Marketplace app connecting dog owners with vetted dog walkers."),
    ("Problem",
     "Urban dog owners work long hours; 40% report skipping walks weekly. "
     "Existing options are informal, uninsured, and unreliable."),
    ("Solution & traction",
     "Vetted walkers, GPS-tracked walks, in-app payments. Live in 3 cities, "
     "1,200 weekly active walkers, $38k MRR growing 15% month over month."),
    ("Business model",
     "20% take rate on each walk. Average walk $25. Subscriptions for "
     "5-walk weekly bundles drive 60% of revenue."),
    ("Team",
     "Two founders: ex-Uber ops lead and a full-stack engineer who built "
     "pet-sitting marketplace Snoutly (acquired 2023)."),
]


def main() -> None:
    prs = Presentation()
    layout = prs.slide_layouts[1]  # Title and Content
    for title, body in SLIDES:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    out = Path(tempfile.gettempdir()) / "rover-sample-deck.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
