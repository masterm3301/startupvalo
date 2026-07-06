import io

import pytest

from src.engine.deck import MAX_DECK_CHARS, DeckError, extract_deck_text

LONG_NOTE = "This startup sells AI-powered dog walking subscriptions to urban pet owners. " * 5


def _make_pptx(slide_texts, notes=None):
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[5]  # "Title Only"
    for i, text in enumerate(slide_texts):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = text
        if notes and i < len(notes) and notes[i]:
            slide.notes_slide.notes_text_frame.text = notes[i]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_pdf(pages):
    from fpdf import FPDF

    pdf = FPDF()
    for text in pages:
        pdf.add_page()
        pdf.set_font("helvetica", size=12)
        pdf.multi_cell(0, 10, text)
    return bytes(pdf.output())


def test_pptx_extracts_slides_and_notes():
    data = _make_pptx(["Problem: dogs need walks", "Market: $1.85B"],
                      notes=[LONG_NOTE, None])
    text = extract_deck_text("pitch.pptx", data)
    assert "Slide 1:" in text and "Slide 2:" in text
    assert "dogs need walks" in text
    assert "urban pet owners" in text  # speaker notes included


def test_pdf_extracts_pages():
    data = _make_pdf([LONG_NOTE, "Traction: 3 cities live"])
    text = extract_deck_text("pitch.pdf", data)
    assert "Page 1:" in text and "Page 2:" in text
    assert "3 cities" in text


def test_legacy_ppt_rejected():
    with pytest.raises(DeckError, match="save as .pptx"):
        extract_deck_text("old-deck.ppt", b"whatever")


def test_unsupported_extension_rejected():
    with pytest.raises(DeckError, match="Unsupported file type"):
        extract_deck_text("pitch.docx", b"whatever")


def test_corrupt_file_rejected():
    with pytest.raises(DeckError, match="corrupt"):
        extract_deck_text("pitch.pptx", b"this is not a zip archive")


def test_near_empty_deck_rejected():
    data = _make_pptx(["Hi"])
    with pytest.raises(DeckError, match="readable text"):
        extract_deck_text("pitch.pptx", data)


def test_long_deck_truncated():
    data = _make_pdf([LONG_NOTE * 60])  # well over 15k chars
    text = extract_deck_text("pitch.pdf", data)
    assert len(text) <= MAX_DECK_CHARS


def test_parser_deck_error_passes_through(monkeypatch):
    from src.engine import deck

    def raising_parser(data):
        raise DeckError("specific parser message")

    monkeypatch.setitem(deck._PARSERS, ".pptx", raising_parser)
    with pytest.raises(DeckError, match="specific parser message"):
        extract_deck_text("pitch.pptx", b"anything")
