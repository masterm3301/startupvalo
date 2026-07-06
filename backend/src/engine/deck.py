"""Extract pitch text from uploaded deck files (.pptx / .pdf)."""
import io
from pathlib import Path

MAX_DECK_CHARS = 15_000
MIN_TEXT_CHARS = 200


class DeckError(ValueError):
    """User-readable deck extraction failure — messages go straight to the UI."""


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    blocks = []
    for n, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"(speaker notes) {notes}")
        if parts:
            blocks.append(f"Slide {n}: " + "\n".join(parts))
    return "\n\n".join(blocks)


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    blocks = []
    for n, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append(f"Page {n}: {text}")
    return "\n\n".join(blocks)


_PARSERS = {".pptx": _pptx_text, ".pdf": _pdf_text}


def extract_deck_text(filename: str, data: bytes) -> str:
    """Return the deck's readable text, or raise DeckError with a user-facing message."""
    ext = Path(filename or "").suffix.lower()
    if ext == ".ppt":
        raise DeckError(
            "Legacy .ppt is not supported — open it in PowerPoint and save as .pptx."
        )
    parser = _PARSERS.get(ext)
    if parser is None:
        raise DeckError(
            f"Unsupported file type {ext or '(none)'} — upload a .pptx or .pdf deck."
        )
    try:
        text = parser(data).strip()
    except Exception:
        raise DeckError(
            f"Could not read {filename} — the file appears corrupt or is not a valid {ext} file."
        )
    if len(text) < MIN_TEXT_CHARS:
        raise DeckError(
            "The deck contains almost no readable text (image-only slides?). "
            "Export a text-based deck or PDF."
        )
    return text[:MAX_DECK_CHARS]
